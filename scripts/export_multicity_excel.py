from pathlib import Path
import csv,sys
from openpyxl import Workbook
from openpyxl.worksheet.table import Table,TableStyleInfo
root=Path(__file__).resolve().parents[1]
source=root/'work/data/aggregates'
out=root/'outputs/LaDe-PowerBI-aggregates.xlsx'
config=[('CityMonth','city_month.csv'),('AcceptanceHour','city_accept_hour.csv'),('RegionMonth','city_region_month.csv')]
wb=Workbook();ws=wb.active;ws.title='ReadMe'
ws.append(['Purpose','Power BI source tables aggregated from the local five-city LaDe-D delivery model.'])
ws.append(['Scope','Historical Chinese delivery data; provisional year 2022; interval is not SLA/on-time/customer satisfaction.'])
ws.append(['Privacy','Contains no order_id or courier_id. Raw source rows are not included.'])
ws.append(['Threshold','Hour and region-month aggregates omit groups under 200 packages.'])
ws.append(['Refresh','Run scripts/build_multicity_models.py and scripts/build_multicity_aggregates.py, then regenerate this workbook.'])
ws.column_dimensions['A'].width=20;ws.column_dimensions['B'].width=115;ws.freeze_panes='A2'
for title,filename in config:
 ws=wb.create_sheet(title)
 with (source/filename).open(encoding='utf-8-sig',newline='') as f:
  for row in csv.reader(f): ws.append(row)
 ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
 for col in ws.columns:
  letter=col[0].column_letter
  ws.column_dimensions[letter].width=min(max(max(len(str(cell.value or '')) for cell in col)+2,12),30)
 if ws.max_row>1:
  ref=f'A1:{ws.cell(ws.max_row,ws.max_column).coordinate}'
  table=Table(displayName=title+'Table',ref=ref)
  table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
  ws.add_table(table)
wb.save(out)
print(f'Created {out} ({out.stat().st_size:,} bytes)')

